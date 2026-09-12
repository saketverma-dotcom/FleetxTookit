import datetime
import time
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import (API_BASE, ASSET_ATTACH_TYPE, ASSET_ATTACH_URL,
                      ASSET_ACCOUNT_URL, ASSET_SUPPLIERS, DEVICE_SUPPLIERS,
                      LOGS_DIR)
from ..api_client import api_headers
from ..http import session
from ..io_utils import load_excel_records
from .. import logic as L


class BulkOnboardTabMixin:
    """Run all five onboarding APIs in order, per row.

        1 Device Add  →  2 Asset Add  →  3 Vehicle-Device Map
        →  4 Asset→Account  →  5 Asset→Vehicle

    assetId is the deviceId (IMEI). If a step fails, the rest of that row is
    skipped so we never attach an asset that was never created; the run
    continues with the next row.
    """

    def _tab_bulk_onboard(self):
        tab = self._scrollable_tab("Bulk Onboard", padding=8)

        ttk.Label(tab, text="Runs 5 APIs in order per row: Device Add → Asset Add → "
                            "Vehicle Map → Asset→Account → Asset→Vehicle",
                  font=("Segoe UI", 9, "bold")).pack(anchor="w")
        ttk.Label(tab, text="A failed step stops that row (later steps are marked SKIPPED); "
                            "other rows keep running.",
                  foreground="gray").pack(anchor="w", pady=(0, 6))

        # ── run-level defaults ──
        d = ttk.LabelFrame(tab, text="Defaults for this run "
                                     "(a matching column in your data overrides these)",
                           padding=6)
        d.pack(fill="x", pady=2)

        g1 = ttk.Frame(d); g1.pack(fill="x", pady=2)
        ttk.Label(g1, text="Device Type:").pack(side="left")
        self.bo_device_type = tk.StringVar(value="LCD40AI-2CH")
        ttk.Combobox(g1, textvariable=self.bo_device_type, width=20,
                     values=["LCD40AI-2CH", "LCD40", "LCD603", "FMB920",
                             "Cello-CANiQ 2G K-Line"]).pack(side="left", padx=4)
        ttk.Label(g1, text="Device Supplier:").pack(side="left", padx=(12, 0))
        self.bo_device_supplier = tk.StringVar(value="CLIENT")
        ttk.Combobox(g1, textvariable=self.bo_device_supplier, width=12,
                     values=DEVICE_SUPPLIERS).pack(side="left", padx=4)
        ttk.Label(g1, text="accountId required on every row",
                  foreground="#c62828").pack(side="left", padx=(12, 0))

        g2 = ttk.Frame(d); g2.pack(fill="x", pady=2)
        ttk.Label(g2, text="Asset Model:").pack(side="left")
        self.bo_asset_model = tk.StringVar(value="LCD40AI-2CH")
        ttk.Entry(g2, textvariable=self.bo_asset_model, width=18).pack(side="left", padx=4)
        ttk.Label(g2, text="Asset Type:").pack(side="left", padx=(10, 0))
        self.bo_asset_type = tk.StringVar(value="CAMERA")
        ttk.Entry(g2, textvariable=self.bo_asset_type, width=14).pack(side="left", padx=4)
        ttk.Label(g2, text="Supplier:").pack(side="left", padx=(10, 0))
        self.bo_supplier = tk.StringVar(value="CLIENT")
        ttk.Combobox(g2, textvariable=self.bo_supplier, width=12,
                     values=ASSET_SUPPLIERS).pack(side="left", padx=4)
        ttk.Label(g2, text="issuedToUserId:").pack(side="left", padx=(10, 0))
        self.bo_user = tk.StringVar(value="14203")
        ttk.Entry(g2, textvariable=self.bo_user, width=10).pack(side="left", padx=4)

        ttk.Label(tab, text="Paste:  deviceId, vehicleId, accountId  (all three required)"
                            "  [, sim [, serialNumber [, deviceType]]]",
                  foreground="gray").pack(anchor="w", pady=(6, 0))
        ttk.Label(tab, text="Skipping sim/serial: omit them, or leave the positions empty "
                            "— e.g.  imei, vehicleId, accountId, , , FMB920",
                  foreground="gray").pack(anchor="w")
        ttk.Label(tab, text="Excel columns:  deviceId | vehicleId | accountId | sim | "
                            "serialNumber | deviceType   (blank cells are fine)",
                  foreground="gray").pack(anchor="w")

        # ── pacing ──
        p = ttk.Frame(tab); p.pack(fill="x", pady=(6, 0))
        ttk.Label(p, text="Gap between steps (ms):").pack(side="left")
        self.bo_step_delay = tk.StringVar(value="600")
        ttk.Entry(p, textvariable=self.bo_step_delay, width=7).pack(side="left", padx=4)
        ttk.Label(p, text="Gap between rows (ms):").pack(side="left", padx=(14, 0))
        self.bo_row_delay = tk.StringVar(value="1000")
        ttk.Entry(p, textvariable=self.bo_row_delay, width=7).pack(side="left", padx=4)
        ttk.Label(p, text="(a gap gives FleetX time to register each change before "
                          "the next call depends on it)",
                  foreground="gray").pack(side="left", padx=6)
        self.bo_src = self._input_source(tab, "Onboarding rows")

        b = ttk.Frame(tab); b.pack(fill="x", pady=2)
        ttk.Button(b, text="▶ Run Bulk Onboard",
                   command=lambda: self._run_thread(self._run_bulk_onboard)
                   ).pack(side="left")
        ttk.Button(b, text="⬇ Sample Excel",
                   command=self._bo_sample).pack(side="left", padx=8)
        self.bo_dry = tk.BooleanVar(value=False)
        ttk.Checkbutton(b, text="Dry run (validate only, no API calls)",
                        variable=self.bo_dry).pack(side="left", padx=12)

    # ── input ──

    def _bo_defaults(self):
        return {"deviceType": self.bo_device_type.get(),
                "deviceSupplier": self.bo_device_supplier.get(),
                "assetModel": self.bo_asset_model.get(),
                "assetType": self.bo_asset_type.get(),
                "assetSupplier": self.bo_supplier.get(),
                "issuedToUserId": self.bo_user.get()}

    def _bo_delays(self):
        """(step_gap_seconds, row_gap_seconds) from the pacing fields."""
        def ms(var, default):
            try:
                v = int(float(var.get().strip() or default))
                return max(0, min(60000, v)) / 1000.0
            except Exception:
                return default / 1000.0
        return ms(self.bo_step_delay, 600), ms(self.bo_row_delay, 1000)

    def _bo_collect(self):
        defaults = self._bo_defaults()
        src = self.bo_src
        if src["mode"].get() == "excel":
            path = src["path"].get().strip()
            if not path:
                self._ui_error("Input", "Select an Excel file.")
                return None
            records = self._load_excel_safe(load_excel_records, path)
            if records is None:
                return None
            rows, errors = L.parse_onboard_records(records, defaults)
        else:
            rows, errors = L.parse_onboard_paste(src["paste"].get("1.0", "end"), defaults)

        for e in errors[:20]:
            self.log(f"  ⚠ {e}", "err")
        if len(errors) > 20:
            self.log(f"  ⚠ ...and {len(errors) - 20} more skipped rows.", "err")
        if not rows:
            self._ui_error("Input", "No valid rows to process."
                           + ("\n\n" + errors[0] if errors else ""))
            return None
        if errors and not self._ui_askyesno(
                "Skip invalid rows?",
                f"{len(errors)} row(s) will be skipped.\n"
                f"Continue with the {len(rows)} valid row(s)?"):
            return None
        return rows

    # ── the five steps ──

    def _bo_step_calls(self, row):
        """(step name, callable) for each step, in order."""
        hdr_json = api_headers(self.token, content_type="application/json")

        def device_add():
            return session.post(f"{API_BASE}/api/v1/devices/",
                                json=L.build_onboard_device(row, self._bo_defaults()),
                                headers=api_headers(self.token), timeout=30)

        def asset_add():
            return session.post(f"{API_BASE}/api/v1/assets",
                                json=L.build_onboard_asset(row, self._bo_defaults()),
                                headers=hdr_json, timeout=30)

        def vehicle_map():
            # multipart form, like the standalone mapping tab
            return session.post(
                f"{API_BASE}/api/v1/vehicles/device",
                files={"deviceId": (None, str(row["device_id"])),
                       "vehicleId": (None, str(row["vehicle_id"]))},
                headers={k: v for k, v in api_headers(self.token).items()
                         if k != "content-type"},
                timeout=30)

        def asset_account():
            params, body = L.build_account_assign(row)
            return session.put(ASSET_ACCOUNT_URL, params=params, json=body,
                               headers=hdr_json, timeout=30)

        def asset_vehicle():
            return session.post(ASSET_ATTACH_URL,
                                json=L.build_attach_body(row, ASSET_ATTACH_TYPE),
                                headers=hdr_json, timeout=30)

        return list(zip(L.ONBOARD_STEPS,
                        [device_add, asset_add, vehicle_map,
                         asset_account, asset_vehicle]))

    # ── run ──

    def _run_bulk_onboard(self):
        rows = self._bo_collect()
        if not rows:
            return
        dry = self.bo_dry.get()

        if dry:
            self.log(f"── DRY RUN — {len(rows)} row(s), no API calls ──", "info")
            for r in rows:
                self.log(f"  {r['device_id']} → vehicle {r['vehicle_id']}, "
                         f"account {r['account_id']}, type {r['device_type']}", "info")
            self.log("Dry run complete.", "ok")
            return

        step_gap, row_gap = self._bo_delays()
        self.log(f"── Bulk Onboard: {len(rows)} row(s) × 5 steps "
                 f"(step gap {step_gap:g}s, row gap {row_gap:g}s) ──", "info")
        results = []
        for idx, row in enumerate(rows, start=1):
            if idx > 1 and row_gap and not self.stop_flag:
                time.sleep(row_gap)
            if self.stop_flag:
                self.log("STOPPED by user.", "err")
                break
            statuses, halted = {}, False
            self.log(f"[{idx}/{len(rows)}] device {row['device_id']}", "info")
            for step_no, (name, call) in enumerate(self._bo_step_calls(row)):
                if halted:
                    statuses[name] = ("SKIPPED", "", "")
                    continue
                if step_no > 0 and step_gap:
                    time.sleep(step_gap)          # let the previous change register
                if self.stop_flag:
                    statuses[name] = ("SKIPPED", "", "stopped by user")
                    halted = True
                    continue
                try:
                    r = call()
                    ok = 200 <= r.status_code < 300
                    body = (r.text or "")[:300]
                    # A 409 on a CREATE step means the object already exists —
                    # e.g. a previous partial run created it. That must not
                    # stop the row, or those rows could never finish steps 3-5.
                    already = (r.status_code == 409
                               and name in ("Device Add", "Asset Add"))
                    if already:
                        statuses[name] = ("EXISTS", r.status_code, body)
                        self.log(f"    ⤼ {name} already exists (409) — continuing",
                                 "info")
                        continue
                    statuses[name] = ("SUCCESS" if ok else "FAILED", r.status_code, body)
                    self.log(f"    {'✓' if ok else '✗'} {name} ({r.status_code})",
                             "ok" if ok else "err")
                    if not ok:
                        halted = True
                        self.log(f"      ↳ stopping this row; remaining steps skipped",
                                 "err")
                        if body:
                            self.log(f"      ↳ {body[:160]}", "err")
                except Exception as e:
                    statuses[name] = ("ERROR", "", str(e)[:300])
                    self.log(f"    ✗ {name} — {e}", "err")
                    halted = True
            results.append((row, statuses, not halted))

        ok_rows = sum(1 for _, _, done in results if done)
        self.log(f"── Done: {ok_rows}/{len(results)} row(s) completed all 5 steps ──",
                 "ok" if ok_rows == len(results) else "err")
        self._bo_save_log(results)

    # ── result workbook: one row per input, a column per step ──

    def _bo_save_log(self, results):
        if not results:
            return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Bulk Onboard"
            headers = (["Device ID", "Asset ID", "Vehicle ID", "Account ID"]
                       + L.ONBOARD_STEPS + ["All Steps OK", "Timestamp"])
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True)
            green = PatternFill("solid", fgColor="C6EFCE")
            red = PatternFill("solid", fgColor="FFC7CE")
            grey = PatternFill("solid", fgColor="E0E0E0")
            yellow = PatternFill("solid", fgColor="FFEB9C")
            ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            for row, statuses, done in results:
                line = [row["device_id"], row["asset_id"], row["vehicle_id"],
                        row["account_id"]]
                for s in L.ONBOARD_STEPS:
                    state, code, _ = statuses.get(s, ("", "", ""))
                    line.append(f"{state} {code}".strip())
                line += ["YES" if done else "NO", ts]
                ws.append(line)
                for i, s in enumerate(L.ONBOARD_STEPS, start=5):
                    cell = ws.cell(row=ws.max_row, column=i)
                    txt = str(cell.value or "")
                    cell.fill = (green if txt.startswith("SUCCESS")
                                 else yellow if txt.startswith("EXISTS")
                                 else grey if txt.startswith("SKIPPED")
                                 else red if txt else None) or cell.fill
            for i in range(1, len(headers) + 1):
                ws.column_dimensions[get_column_letter(i)].width = 18

            import os
            os.makedirs(LOGS_DIR, exist_ok=True)
            path = os.path.join(
                LOGS_DIR,
                f"bulk_onboard_{datetime.datetime.now():%Y%m%d_%H%M%S}.xlsx")
            wb.save(path)
            self.log(f"Result log saved: {path}", "ok")
        except Exception as e:
            self.log(f"Could not save result log: {e}", "err")

    def _bo_sample(self):
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialfile="bulk_onboard_sample.xlsx",
            filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Input"
            ws.append(["deviceId", "vehicleId", "accountId", "sim",
                       "serialNumber", "deviceType"])
            ws.append(["862798051206510", 2793621, 17809, "8991000012345",
                       "SN123", "LCD40AI-2CH"])
            ws.append(["862798051206511", 2793622, 17809, "", "", ""])
            for c in ws[1]:
                c.font = Font(bold=True)
            for i in range(1, 7):
                ws.column_dimensions[get_column_letter(i)].width = 20
            wb.save(path)
            messagebox.showinfo(
                "Sample Excel",
                f"Saved:\n{path}\n\naccountId is required on every row. "
                f"A blank deviceType falls back to the tab default.")
        except Exception as e:
            self._ui_error("Sample Excel", f"Could not save:\n{e}")
