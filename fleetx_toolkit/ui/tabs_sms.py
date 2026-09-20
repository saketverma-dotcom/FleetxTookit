import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

import requests
from ..http import session

from ..config import (LOGS_DIR, SEMYSMS_API, SEMYSMS_SIM_NAMES,
                      sim_id_for_name, sim_name_for_id)
from .. import messaging as _M
from ..io_utils import load_excel_records
from ..sms import build_sms_params, normalize_number, sms_success
from ..storage import load_sms_token, save_sms_token


class SmsTabMixin:
    """SMS Command tab — bulk SMS via SemySMS, per-row SIM selection by name."""

    def _tab_sms(self):
        tab = self._scrollable_tab("SMS Command", padding=8)

        # ── Token row (stored in Windows Credential Manager) ──
        tf = ttk.LabelFrame(tab, text="SemySMS token (stored securely on this PC)", padding=6)
        tf.pack(fill="x")
        self.sms_token = tk.StringVar(value=self._current_sms_token())
        ttk.Entry(tf, textvariable=self.sms_token, width=40, show="•").pack(side="left", padx=4)
        ttk.Button(tf, text="Save token", command=self._save_sms_token).pack(side="left", padx=4)
        self.sms_token_status = ttk.Label(tf, text="", foreground="green")
        self.sms_token_status.pack(side="left", padx=6)

        # ── SIM selector (tab default) ──
        sf = ttk.Frame(tab); sf.pack(fill="x", pady=(8, 2))
        ttk.Label(sf, text="SIM to send from:").pack(side="left")
        self.sms_sim = tk.StringVar(value=SEMYSMS_SIM_NAMES[0])
        _sim_cb = ttk.Combobox(sf, textvariable=self.sms_sim, width=28,
                               state="readonly", values=SEMYSMS_SIM_NAMES)
        _sim_cb.pack(side="left", padx=6)
        # refresh battery/online when opened (cached 60s, fetched off-thread)
        _sim_cb.bind("<Button-1>", lambda e: self._sim_refresh_status())
        _sim_cb.bind("<<ComboboxSelected>>", lambda e: self._sim_apply_status())
        self._sms_sim_cb = _sim_cb
        self.sms_sim_state = ttk.Label(sf, text="", font=("Segoe UI", 9, "bold"))
        self.sms_sim_state.pack(side="left")
        ttk.Label(sf, text="(used for pasted numbers, and for Excel rows with a blank SIM Name)",
                  foreground="gray").pack(side="left")

        # ── Mode selector ──
        self.sms_mode = tk.StringVar(value="paste")
        mf = ttk.Frame(tab); mf.pack(fill="x", pady=(8, 2))
        ttk.Radiobutton(mf, text="Same message → many numbers", value="paste",
                        variable=self.sms_mode, command=self._sms_mode_switch).pack(side="left")
        ttk.Radiobutton(mf, text="Per-row from Excel (Mobile | Message | SIM Name)", value="excel",
                        variable=self.sms_mode, command=self._sms_mode_switch).pack(side="left", padx=12)

        # ── Paste frame ──
        self.sms_paste_frame = ttk.Frame(tab)
        nrow = ttk.Frame(self.sms_paste_frame); nrow.pack(fill="x")
        ttk.Label(nrow, text="Numbers (one per line):").pack(side="left")
        self.sms_count_lbl = ttk.Label(nrow, text="0 recipients", foreground="gray")
        self.sms_count_lbl.pack(side="right")
        self.sms_numbers = scrolledtext.ScrolledText(self.sms_paste_frame, height=6, width=40)
        self.sms_numbers.pack(fill="x")
        self.sms_numbers.bind("<KeyRelease>", lambda e: self._sms_update_counts())
        mrow = ttk.Frame(self.sms_paste_frame); mrow.pack(fill="x", pady=(6, 0))
        ttk.Label(mrow, text="Message:").pack(side="left")
        self.sms_char_lbl = ttk.Label(mrow, text="0 chars · 0 SMS", foreground="gray")
        self.sms_char_lbl.pack(side="right")
        self.sms_msg = scrolledtext.ScrolledText(self.sms_paste_frame, height=3, width=40)
        self.sms_msg.pack(fill="x")
        self.sms_msg.bind("<KeyRelease>", lambda e: self._sms_update_counts())

        # ── Excel frame ──
        self.sms_excel_frame = ttk.Frame(tab)
        er = ttk.Frame(self.sms_excel_frame); er.pack(fill="x")
        self.sms_path = tk.StringVar()
        ttk.Entry(er, textvariable=self.sms_path, width=52).pack(side="left", padx=4)
        ttk.Button(er, text="Browse...",
                   command=lambda: self.sms_path.set(
                       filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")]))).pack(side="left")
        ttk.Button(er, text="⬇ Sample Excel",
                   command=lambda: self._download_sample("sms", "sms_sample.xlsx")).pack(side="left", padx=4)

        self._sms_mode_switch()

        ttk.Label(tab, text="Country code to prepend (blank = send as-is):",
                  foreground="gray").pack(anchor="w", pady=(8, 0))
        self.sms_cc = tk.StringVar(value="")
        ttk.Entry(tab, textvariable=self.sms_cc, width=8).pack(anchor="w")

        dr = ttk.Frame(tab); dr.pack(anchor="w", pady=(8, 0))
        ttk.Label(dr, text="Delay between SMS (seconds):").pack(side="left")
        self.sms_delay = tk.StringVar(value="5")
        ttk.Entry(dr, textvariable=self.sms_delay, width=6).pack(side="left", padx=6)
        ttk.Label(dr, text="(overrides the global Settings delay for SMS only)",
                  foreground="gray").pack(side="left")

        # ── delivery report controls ──
        dl = ttk.Frame(tab); dl.pack(anchor="w", pady=(6, 0))
        self.sms_track = tk.BooleanVar(value=True)
        ttk.Checkbutton(dl, text="Track delivery status", variable=self.sms_track,
                        command=self._sms_toggle_tracking).pack(side="left")
        ttk.Label(dl, text="Give up after (min):").pack(side="left", padx=(14, 0))
        self.sms_track_timeout = tk.StringVar(value="2")
        ttk.Entry(dl, textvariable=self.sms_track_timeout, width=5).pack(side="left", padx=4)
        ttk.Label(dl, text="poll every (s):").pack(side="left", padx=(10, 0))
        self.sms_track_poll = tk.StringVar(value="10")
        ttk.Entry(dl, textvariable=self.sms_track_poll, width=5).pack(side="left", padx=4)
        ttk.Label(dl, text="— messages still pending at the timeout are logged as Pending",
                  foreground="gray").pack(side="left", padx=6)

        ttk.Button(tab, text="▶ Send SMS",
                   command=lambda: self._run_thread(self._run_sms)).pack(anchor="w", pady=8)

        # ── live status table (Treeview: native, no per-row widgets to rebuild) ──
        self.sms_status_frame = ttk.LabelFrame(tab, text="Delivery status", padding=4)
        self.sms_status_frame.pack(fill="both", expand=True, pady=(4, 0))
        cols = ("num", "mobile", "sim", "sent", "status", "updated")
        self.sms_tree = ttk.Treeview(self.sms_status_frame, columns=cols,
                                     show="headings", height=9)
        for c, t, w in [("num","#",40),("mobile","Mobile",130),("sim","SIM",120),
                        ("sent","Sent",70),("status","Delivery status",150),
                        ("updated","Updated",90)]:
            self.sms_tree.heading(c, text=t)
            self.sms_tree.column(c, width=w, anchor="w")
        vs=ttk.Scrollbar(self.sms_status_frame, orient="vertical",
                         command=self.sms_tree.yview)
        self.sms_tree.configure(yscrollcommand=vs.set)
        self.sms_tree.pack(side="left", fill="both", expand=True)
        vs.pack(side="right", fill="y")
        self.sms_tree.tag_configure("ok", foreground="#1b7f3b")
        self.sms_tree.tag_configure("bad", foreground="#c62828")
        self.sms_tree.tag_configure("wait", foreground="#e65100")
        self.sms_summary = ttk.Label(tab, text="", foreground="gray")
        self.sms_summary.pack(anchor="w")

    def _sms_toggle_tracking(self):
        """Hide the status table when tracking is switched off."""
        if not hasattr(self, "sms_status_frame"):
            return
        if self.sms_track.get():
            if not self.sms_status_frame.winfo_ismapped():
                self.sms_status_frame.pack(fill="both", expand=True, pady=(4, 0))
        else:
            self.sms_status_frame.pack_forget()

    def _sms_update_counts(self):
        from .. import messaging as _M
        nums = [n for n in (x.strip() for x in
                self.sms_numbers.get("1.0", "end").splitlines()) if n]
        self.sms_count_lbl.config(text=f"{len(nums)} recipient" + ("" if len(nums) == 1 else "s"))
        chars, seg = _M.sms_segments(self.sms_msg.get("1.0", "end").rstrip("\n"))
        self.sms_char_lbl.config(text=f"{chars} chars · {seg} SMS")

    def _sms_mode_switch(self):
        if self.sms_mode.get() == "paste":
            self.sms_excel_frame.pack_forget()
            self.sms_paste_frame.pack(fill="x", pady=4)
        else:
            self.sms_paste_frame.pack_forget()
            self.sms_excel_frame.pack(fill="x", pady=4)

    def _save_sms_token(self):
        tok = self.sms_token.get().strip()
        if not tok:
            self.sms_token_status.config(text="Enter a token first.", foreground="red")
            return
        if save_sms_token(tok):
            self.sms_token_status.config(text="Saved securely.", foreground="green")
        else:
            self.sms_token_status.config(
                text="Could not store securely (session only).", foreground="red")

    def _sms_rows(self):
        """Build [(phone, msg, sim_id), ...] from the active mode, or None on error.
        Applies number normalization and per-row SIM name -> id resolution."""
        cc = self.sms_cc.get().strip()
        default_sim = sim_id_for_name(self.sms_sim.get())
        rows = []
        if self.sms_mode.get() == "paste":
            msg = self.sms_msg.get("1.0", "end").strip()
            if not msg:
                self._ui_error("SMS", "Enter a message."); return None
            nums = [n for n in (normalize_number(x, cc)
                                for x in self.sms_numbers.get("1.0", "end").splitlines()) if n]
            if not nums:
                self._ui_error("SMS", "Enter at least one number."); return None
            rows = [(n, msg, default_sim) for n in nums]
        else:
            path = self.sms_path.get().strip()
            if not path:
                self._ui_error("SMS", "Select an Excel file."); return None
            records = self._load_excel_safe(load_excel_records, path)
            if records is None:
                return None
            for rec in records:
                phone = normalize_number(rec.get("mobile", ""), cc)
                msg = str(rec.get("message", "") or "").strip()
                raw_sim = rec.get("sim_name", "")
                sim_name = "" if raw_sim is None else str(raw_sim).strip()
                if sim_name.lower() == "none":
                    sim_name = ""
                sim_id = sim_id_for_name(sim_name) if sim_name else default_sim
                if not phone or not msg:
                    continue
                if not sim_id:
                    self._ui_error("SMS",
                        f"Row for {phone}: unknown SIM Name '{sim_name}'.\n"
                        f"Use one of: {', '.join(SEMYSMS_SIM_NAMES)}")
                    return None
                rows.append((phone, msg, sim_id))
            if not rows:
                self._ui_error("SMS", "No valid rows (need Mobile + Message)."); return None
        return rows

    def _run_sms(self):
        tok = self.sms_token.get().strip() or self._current_sms_token()
        if not tok:
            self._sms_log("  ✗ No SemySMS token — enter and Save it first.", "err")
            self._ui_error("SMS", "Enter and save the SemySMS token first.")
            return
        rows = self._sms_rows()
        if not rows:
            return

        # Dedicated SMS delay (seconds) overrides the global Settings delay for
        # this run only; restored in finally so other tabs are unaffected.
        try:
            sms_delay = float(self.sms_delay.get().strip() or 5)
            if sms_delay < 0:
                raise ValueError
        except ValueError:
            self._sms_log(f"  ✗ Invalid SMS delay '{self.sms_delay.get()}' — must be seconds ≥ 0.", "err")
            self._ui_error("SMS", "Delay between SMS must be a number of seconds (e.g. 5).")
            return

        # Delivery tracking needs the SemySMS message id, which the send
        # response carries — the old code checked only code=="0" and dropped it.
        track = bool(getattr(self, "sms_track", None) and self.sms_track.get())
        self._sms_sent = []          # [{n, phone, msg, sim_id, sms_id, sent_ok}]
        if track:
            self.after(0, self._sms_reset_table)

        def fn(row):
            phone, msg, sim_id = row
            r = session.post(SEMYSMS_API,
                              data=build_sms_params(tok, sim_id, phone, msg), timeout=30)
            ok, sms_id = False, 0
            try:
                body = r.json() or {}
                ok = sms_success(body)
                sms_id = int(body.get("id") or 0)
            except Exception:
                ok = False
            if r.status_code == 200 and not ok:
                r.status_code = 422   # surface API rejection as a failed row
            if track:
                rec = {"n": len(self._sms_sent)+1, "phone": phone, "msg": msg,
                       "sim_id": sim_id, "sms_id": sms_id, "sent_ok": ok,
                       "status": _M.SENT_PENDING if ok else _M.SENT_FAILED,
                       "updated": ""}
                self._sms_sent.append(rec)
                self.after(0, lambda r=rec: self._sms_row_upsert(r))
            return (phone, msg[:40], sim_id), r

        self._delay_override = sms_delay
        try:
            # LOG 1 — the send result, saved by _loop as it always was
            self._loop(rows, "SMS Send", fn, ["Mobile", "Message", "SIM ID"])
        finally:
            self._delay_override = None

        if track:
            self._sms_track_delivery()

    # ─────────── delivery status tracking (v3.17) ───────────

    def _sms_reset_table(self):
        if not hasattr(self, "sms_tree"):
            return
        for i in self.sms_tree.get_children():
            self.sms_tree.delete(i)
        self.sms_summary.config(text="")

    def _sms_row_upsert(self, rec):
        """Insert or update ONE row in place. Treeview is a native widget, so
        this never rebuilds the table — the lag trap from the Messaging list."""
        if not hasattr(self, "sms_tree"):
            return
        iid = f"r{rec['n']}"
        label = _M.STATUS_LABEL.get(rec["status"], rec["status"])
        tag = ("ok" if rec["status"] == _M.SENT_DELIVERED else
               "bad" if rec["status"] in (_M.SENT_FAILED, _M.SENT_CANCELLED) else "wait")
        vals = (rec["n"], rec["phone"], sim_name_for_id(rec["sim_id"]) or rec["sim_id"],
                "✓" if rec["sent_ok"] else "✗", label, rec.get("updated", ""))
        try:
            if self.sms_tree.exists(iid):
                self.sms_tree.item(iid, values=vals, tags=(tag,))
            else:
                self.sms_tree.insert("", "end", iid=iid, values=vals, tags=(tag,))
                self.sms_tree.see(iid)
        except Exception:
            pass

    def _sms_track_settings(self):
        def num(var, default, lo, hi):
            try:
                v = float(str(var.get()).strip() or default)
                return max(lo, min(hi, v))
            except Exception:
                return default
        return (num(self.sms_track_timeout, 2, 0.25, 60) * 60.0,   # seconds
                num(self.sms_track_poll, 10, 3, 120))

    def _sms_track_delivery(self):
        """Poll outbox_sms.php until every message reaches a terminal state or
        the timeout expires. Ids are batched per SIM, so this costs a handful of
        calls per cycle rather than one per message."""
        import time
        pending = [r for r in self._sms_sent if r["sent_ok"] and r["sms_id"]]
        if not pending:
            self._sms_log("No message ids returned — delivery tracking skipped.", "err")
            self._sms_finalise_delivery(timed_out=False)
            return
        timeout_s, poll_s = self._sms_track_settings()
        tok = self._current_sms_token()
        self._sms_log(f"Tracking delivery for {len(pending)} message(s) — "
                      f"polling every {poll_s:g}s, giving up after {timeout_s/60:g} min.", "info")
        by_id = {r["sms_id"]: r for r in pending}
        deadline = time.monotonic() + timeout_s
        timed_out = False
        while True:
            open_ids = [i for i, r in by_id.items()
                        if r["status"] not in _M.TERMINAL_STATUSES]
            if not open_ids:
                break
            if self.stop_flag:
                self._sms_log("STOP pressed — finalising with the statuses known so far.", "err")
                break
            if time.monotonic() >= deadline:
                timed_out = True
                self._sms_log(f"Delivery timeout reached — {len(open_ids)} still pending.", "err")
                break
            # group by SIM (outbox is per-device), then batch the ids
            groups = {}
            for i in open_ids:
                groups.setdefault(by_id[i]["sim_id"], []).append(i)
            for device, ids in groups.items():
                for k in range(0, len(ids), 50):          # conservative batch
                    if self.stop_flag:
                        break
                    chunk = ids[k:k+50]
                    url, params = _M.build_outbox_request(tok, device, chunk)
                    try:
                        r = session.get(url, params=params, timeout=30)
                        statuses = _M.parse_outbox_status(r.json())
                    except Exception as e:
                        self._sms_log(f"  status poll error: {_M.friendly_error(e)}", "err")
                        statuses = {}
                    stamp = __import__("datetime").datetime.now().strftime("%H:%M:%S")
                    for sid, st in statuses.items():
                        rec = by_id.get(sid)
                        if rec and rec["status"] != st:
                            rec["status"] = st
                            rec["updated"] = stamp
                            self.after(0, lambda r=rec: self._sms_row_upsert(r))
            self.after(0, self._sms_update_summary)
            for _ in range(int(poll_s * 2)):
                if self.stop_flag:
                    break
                time.sleep(0.5)
        self.after(0, self._sms_update_summary)
        self._sms_finalise_delivery(timed_out=timed_out)

    def _sms_update_summary(self):
        if not hasattr(self, "sms_summary"):
            return
        c = {}
        for r in getattr(self, "_sms_sent", []):
            c[r["status"]] = c.get(r["status"], 0) + 1
        parts = [f"{_M.STATUS_LABEL.get(k,k)}: {v}" for k, v in sorted(c.items())]
        self.sms_summary.config(text="   ".join(parts))

    def _sms_finalise_delivery(self, timed_out):
        """LOG 2 — the delivery report, auto-saved to the logs folder."""
        import datetime, os
        rows = getattr(self, "_sms_sent", [])
        if not rows:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SMS Delivery"
            hdr = ["#", "Mobile", "SIM", "SemySMS ID", "Message",
                   "Send accepted", "Delivery status", "Last update"]
            ws.append(hdr)
            for c in ws[1]:
                c.font = Font(bold=True)
            green = PatternFill("solid", fgColor="C6EFCE")
            red = PatternFill("solid", fgColor="FFC7CE")
            amber = PatternFill("solid", fgColor="FFEB9C")
            for r in rows:
                ws.append([r["n"], r["phone"],
                           sim_name_for_id(r["sim_id"]) or r["sim_id"],
                           r["sms_id"] or "", r["msg"][:120],
                           "YES" if r["sent_ok"] else "NO",
                           _M.STATUS_LABEL.get(r["status"], r["status"]),
                           r.get("updated", "")])
                cell = ws.cell(row=ws.max_row, column=7)
                cell.fill = (green if r["status"] == _M.SENT_DELIVERED else
                             red if r["status"] in (_M.SENT_FAILED, _M.SENT_CANCELLED)
                             else amber)
            for i in range(1, len(hdr) + 1):
                ws.column_dimensions[get_column_letter(i)].width = 18
            ws.column_dimensions["E"].width = 42
            os.makedirs(LOGS_DIR, exist_ok=True)
            path = os.path.join(
                LOGS_DIR, f"sms_delivery_{datetime.datetime.now():%Y%m%d_%H%M%S}.xlsx")
            wb.save(path)
            done = sum(1 for r in rows if r["status"] == _M.SENT_DELIVERED)
            self._sms_log(f"Delivery report saved ({done}/{len(rows)} delivered"
                          + (", timed out" if timed_out else "") + f"): {path}", "ok")
        except Exception as e:
            self._sms_log(f"Could not save the delivery report: {e}", "err")
