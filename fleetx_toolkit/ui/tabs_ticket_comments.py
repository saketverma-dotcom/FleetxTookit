import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..api_client import api_headers
from ..config import API_BASE
from ..http import session
from ..io_utils import load_excel_records
from .. import logic as L


class TicketCommentsTabMixin:
    """Bulk-comment on internal tickets, optionally re-assigning afterwards.

        POST /api/v1/internal/comment/              one comment per ticket
        PUT  /api/v1/internal/issue/change-assignment   optional, after success

    Two input modes:
      • Paste ticket ids + one comment  → the same comment on every ticket
      • Excel  ticketId | comment | assignTo → a different comment per ticket

    Re-assignment is skipped for any ticket whose comment failed, so a ticket
    is never moved to someone else without the note that explains why.
    """

    def _tab_ticket_comments(self):
        tab = self._scrollable_tab("Ticket Comments", padding=8)

        ttk.Label(tab, text="Comment on many tickets at once, then optionally re-assign them.",
                  font=("Segoe UI", 9, "bold")).pack(anchor="w")
        ttk.Label(tab, text="Admin only. A ticket whose comment fails is never re-assigned.",
                  foreground="gray").pack(anchor="w", pady=(0, 6))

        # ── mode ──
        self.tc_mode = tk.StringVar(value="paste")
        mf = ttk.Frame(tab); mf.pack(anchor="w", pady=2)
        ttk.Radiobutton(mf, text="One comment → many tickets", value="paste",
                        variable=self.tc_mode, command=self._tc_mode_switch).pack(side="left")
        ttk.Radiobutton(mf, text="Excel: a comment per ticket", value="excel",
                        variable=self.tc_mode, command=self._tc_mode_switch).pack(side="left", padx=14)

        # ── paste mode ──
        self.tc_paste_frame = ttk.Frame(tab)
        ttk.Label(self.tc_paste_frame, text="Ticket IDs (comma, space or one per line):"
                  ).pack(anchor="w")
        self.tc_ids = tk.Text(self.tc_paste_frame, height=4, width=60, wrap="word")
        self.tc_ids.pack(fill="x")
        row = ttk.Frame(self.tc_paste_frame); row.pack(fill="x", pady=(4, 0))
        ttk.Label(row, text="Comment:").pack(side="left")
        self.tc_count_lbl = ttk.Label(row, text="0 tickets", foreground="gray")
        self.tc_count_lbl.pack(side="right")
        self.tc_comment = tk.Text(self.tc_paste_frame, height=6, width=60, wrap="word")
        self.tc_comment.pack(fill="x")
        ttk.Label(self.tc_paste_frame,
                  text="Plain text is fine — each line becomes a paragraph. "
                       "Paste HTML and it is sent as-is.",
                  foreground="gray").pack(anchor="w")
        self.tc_ids.bind("<KeyRelease>", lambda e: self._tc_update_count())

        # ── excel mode ──
        self.tc_excel_frame = ttk.Frame(tab)
        ttk.Label(self.tc_excel_frame,
                  text="Columns:  ticketId | comment | assignTo (assignTo optional)"
                  ).pack(anchor="w")
        er = ttk.Frame(self.tc_excel_frame); er.pack(fill="x", pady=2)
        self.tc_path = tk.StringVar()
        ttk.Entry(er, textvariable=self.tc_path, width=52).pack(side="left")
        ttk.Button(er, text="Browse", command=self._tc_browse).pack(side="left", padx=4)
        ttk.Button(er, text="⬇ Sample", command=self._tc_sample).pack(side="left")

        self._tc_mode_switch()

        # ── reassignment ──
        rf = ttk.LabelFrame(tab, text="Re-assignment", padding=6)
        rf.pack(fill="x", pady=8)
        self.tc_reassign = tk.BooleanVar(value=False)
        ttk.Checkbutton(rf, text="Re-assign after a successful comment",
                        variable=self.tc_reassign,
                        command=self._tc_toggle_reassign).pack(anchor="w")
        ar = ttk.Frame(rf); ar.pack(anchor="w", pady=(4, 0))
        ttk.Label(ar, text="Assign to user ID:").pack(side="left")
        self.tc_assignee = tk.StringVar()
        self.tc_assignee_entry = ttk.Entry(ar, textvariable=self.tc_assignee, width=14)
        self.tc_assignee_entry.pack(side="left", padx=6)
        ttk.Label(ar, text="(Excel rows may override this with their own assignTo)",
                  foreground="gray").pack(side="left")
        self._tc_toggle_reassign()

        b = ttk.Frame(tab); b.pack(fill="x", pady=6)
        ttk.Button(b, text="▶ Post Comments",
                   command=lambda: self._run_thread(self._run_ticket_comments)
                   ).pack(side="left")
        self.tc_dry = tk.BooleanVar(value=False)
        ttk.Checkbutton(b, text="Dry run (validate only, no API calls)",
                        variable=self.tc_dry).pack(side="left", padx=12)

    # ── small UI helpers ──

    def _tc_mode_switch(self):
        if self.tc_mode.get() == "excel":
            self.tc_paste_frame.pack_forget()
            self.tc_excel_frame.pack(fill="x", pady=4)
        else:
            self.tc_excel_frame.pack_forget()
            self.tc_paste_frame.pack(fill="x", pady=4)

    def _tc_toggle_reassign(self):
        state = "normal" if self.tc_reassign.get() else "disabled"
        try:
            self.tc_assignee_entry.config(state=state)
        except Exception:
            pass

    def _tc_update_count(self):
        ids, _ = L.parse_ticket_ids(self.tc_ids.get("1.0", "end"))
        self.tc_count_lbl.config(text=f"{len(ids)} ticket" + ("" if len(ids) == 1 else "s"))

    def _tc_browse(self):
        p = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm")])
        if p:
            self.tc_path.set(p)

    def _tc_sample(self):
        p = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                         initialfile="ticket_comments_sample.xlsx",
                                         filetypes=[("Excel", "*.xlsx")])
        if not p:
            return
        try:
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Comments"
            ws.append(["ticketId", "comment", "assignTo"])
            ws.append([407582, "Please align a technician to inspect the GNSS antenna.", 880971])
            ws.append([407583, "Antenna re-seated and verified on site.", ""])
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.column_dimensions["A"].width = 12
            ws.column_dimensions["B"].width = 60
            ws.column_dimensions["C"].width = 14
            wb.save(p)
            messagebox.showinfo("Sample", f"Saved:\n{p}\n\nA blank assignTo falls back to the "
                                          f"user ID on the tab.")
        except Exception as e:
            self._ui_error("Sample", f"Could not save:\n{e}")

    # ── collect ──

    def _tc_collect(self):
        default_assignee = self.tc_assignee.get().strip() if self.tc_reassign.get() else ""
        if self.tc_reassign.get():
            if not default_assignee and self.tc_mode.get() == "paste":
                self._ui_error("Re-assignment", "Enter the user ID to assign to, "
                                                "or switch re-assignment off.")
                return None
            if default_assignee and not default_assignee.isdigit():
                self._ui_error("Re-assignment",
                               f"'{default_assignee}' is not a numeric user ID.")
                return None

        if self.tc_mode.get() == "excel":
            path = self.tc_path.get().strip()
            if not path:
                self._ui_error("Input", "Select an Excel file.")
                return None
            records = self._load_excel_safe(load_excel_records, path)
            if records is None:
                return None
            rows, errors = L.parse_comment_records(records, default_assignee)
        else:
            ids, errors = L.parse_ticket_ids(self.tc_ids.get("1.0", "end"))
            body = self.tc_comment.get("1.0", "end").strip()
            if not body:
                self._ui_error("Input", "Enter the comment to post.")
                return None
            rows = [{"ticket_id": t, "comment": body, "assignee": default_assignee}
                    for t in ids]

        for e in errors[:20]:
            self.log(f"  ⚠ {e}", "err")
        if len(errors) > 20:
            self.log(f"  ⚠ ...and {len(errors)-20} more skipped rows.", "err")
        if not rows:
            self._ui_error("Input", "No valid tickets to process."
                           + ("\n\n" + errors[0] if errors else ""))
            return None
        if errors and not self._ui_askyesno(
                "Skip invalid rows?",
                f"{len(errors)} row(s) will be skipped.\n"
                f"Continue with the {len(rows)} valid one(s)?"):
            return None
        return rows

    # ── run ──

    def _run_ticket_comments(self):
        rows = self._tc_collect()
        if not rows:
            return
        reassign = self.tc_reassign.get()

        if self.tc_dry.get():
            self.log(f"── DRY RUN — {len(rows)} ticket(s), no API calls ──", "info")
            for r in rows[:50]:
                self.log(f"  #{r['ticket_id']}  “{r['comment'][:50]}”"
                         + (f"  → assign {r['assignee']}" if reassign and r['assignee'] else ""),
                         "info")
            self.log("Dry run complete.", "ok")
            return

        if not self._ui_askyesno(
                "Confirm",
                f"Post a comment on {len(rows)} ticket(s)?"
                + (f"\n\nThey will then be re-assigned." if reassign else "")):
            return

        self.log(f"── Ticket comments: {len(rows)} ticket(s) ──", "info")
        results = []
        for i, row in enumerate(rows, start=1):
            if self.stop_flag:
                self.log("STOPPED by user.", "err")
                break
            tid = row["ticket_id"]
            html = L.comment_to_html(row["comment"])
            now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            c_state, c_code, a_state = "FAILED", "", "SKIPPED"
            try:
                r = session.post(f"{API_BASE}/api/v1/internal/comment/",
                                 json=L.build_comment_body(tid, html, now),
                                 headers=api_headers(self.token,
                                                     content_type="application/json"),
                                 timeout=30)
                c_code = r.status_code
                ok = 200 <= r.status_code < 300
                c_state = "SUCCESS" if ok else "FAILED"
                self.log(f"[{i}/{len(rows)}] #{tid} comment {'✓' if ok else '✗'} ({c_code})",
                         "ok" if ok else "err")
                if not ok:
                    self.log(f"      ↳ {(r.text or '')[:160]}", "err")
                # Re-assign only when the comment landed.
                if ok and reassign and row["assignee"]:
                    try:
                        r2 = session.put(
                            f"{API_BASE}/api/v1/internal/issue/change-assignment",
                            json=L.build_reassign_body(tid, row["assignee"]),
                            headers=api_headers(self.token), timeout=30)
                        ok2 = 200 <= r2.status_code < 300
                        a_state = ("SUCCESS" if ok2 else "FAILED") + f" {r2.status_code}"
                        self.log(f"      ↳ re-assign to {row['assignee']} "
                                 f"{'✓' if ok2 else '✗'} ({r2.status_code})",
                                 "ok" if ok2 else "err")
                    except Exception as e:
                        a_state = "ERROR"
                        self.log(f"      ↳ re-assign error: {e}", "err")
                elif reassign and not row["assignee"]:
                    a_state = "NO ASSIGNEE"
            except Exception as e:
                c_state = "ERROR"
                self.log(f"[{i}/{len(rows)}] #{tid} — {e}", "err")
            results.append([tid, row["comment"][:120], f"{c_state} {c_code}".strip(),
                            row["assignee"] or "", a_state, now])

        ok_n = sum(1 for r in results if str(r[2]).startswith("SUCCESS"))
        self.log(f"── Done: {ok_n}/{len(results)} comment(s) posted ──",
                 "ok" if ok_n == len(results) else "err")
        self._tc_save_log(results, reassign)

    def _tc_save_log(self, results, reassign):
        if not results:
            return
        import os
        from ..config import LOGS_DIR
        try:
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Ticket Comments"
            ws.append(["Ticket ID", "Comment", "Comment result",
                       "Assign to", "Re-assign result", "Posted at"])
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in results:
                ws.append(r)
            for i, w in enumerate([12, 60, 18, 12, 18, 20], start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            os.makedirs(LOGS_DIR, exist_ok=True)
            path = os.path.join(
                LOGS_DIR,
                f"ticket_comments_{datetime.datetime.now():%Y%m%d_%H%M%S}.xlsx")
            wb.save(path)
            self.log(f"Result log saved: {path}", "ok")
        except Exception as e:
            self.log(f"Could not save the result log: {e}", "err")
